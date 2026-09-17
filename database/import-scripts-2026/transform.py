import openpyxl, re, json, sys, datetime, collections, os
SRC='/Users/iceth/Desktop/Face Advance Systems/database/2026.xlsx'
FIX='/Users/iceth/Desktop/Face Advance Systems/database/2026-ตรวจก่อนนำเข้า.xlsx'
MONTHS={'ม.ค.69':('1-31 ม.ค.69','2601'),'ก.พ.69':('1-28 ก.พ.69','2602'),'มี.ค.69':('1-31 มี.ค.69','2603'),
        'เม.ย.69':('1-30 เม.ย.69','2604'),'พ.ค.69':('1-31 พ.ค.69','2605'),'มิ.ย.69':('1-30 มิ.ย.69','2606')}
HOPEFUL={"Abbie (Beta)","Lab Farm","Beta Teeth","Beta M","Lab Balance (10 ซอง)","Beta Care Cream","Beta Vit (30 แคปซูล)","Beta Oil กล่อง (10 แคปซูล)","Beta Oil ซอง (5 แคปซูล)","Beta E","Beta Cal Plus กล่อง (10 แคปซูล)","Beta Cal Pro Plus (10 เเคปซูล)","Beta Cal Pro Plus (5 เเคปซูล)","Beta X","Beta X Spray","Beta X Plus","Beta X Plus (5 แคปซูล)","Beta ชากระชายขาว X19","Beta Herb (10 เเคปซูล)","Beta Herb (5 เเคปซูล)","Beta Herb Tea","Beta Life Balance (10 ซอง)","Beta Soap","Beta Flow","Beta Coffee","Memora (10 แคปซูล)","Freshy (Beta)","Hopeful กล่องยา","Hopeful แก้วน้ำ","Hopeful ผ้าเช็ดตัว","Hopeful หมอน","แถมฟรี ปฏิทิน Hopeful","Beta Life (10 เเคปซูล)","Beta Life (5 เเคปซูล)","Beta Liv","Beta Liv Pro Plus (10 แคปซูล)","Beta Liv Pro Plus (5 แคปซูล)","Beta Care (10 แคปซูล)","Hopeful ร่ม"}
PROV=set("กระบี่|กรุงเทพมหานคร|กาญจนบุรี|กาฬสินธุ์|กำแพงเพชร|ขอนแก่น|จันทบุรี|ฉะเชิงเทรา|ชลบุรี|ชัยนาท|ชัยภูมิ|ชุมพร|เชียงราย|เชียงใหม่|ตรัง|ตราด|ตาก|นครนายก|นครปฐม|นครพนม|นครราชสีมา|นครศรีธรรมราช|นครสวรรค์|นนทบุรี|นราธิวาส|น่าน|บึงกาฬ|บุรีรัมย์|ปทุมธานี|ประจวบคีรีขันธ์|ปราจีนบุรี|ปัตตานี|พระนครศรีอยุธยา|พะเยา|พังงา|พัทลุง|พิจิตร|พิษณุโลก|เพชรบุรี|เพชรบูรณ์|แพร่|ภูเก็ต|มหาสารคาม|มุกดาหาร|แม่ฮ่องสอน|ยโสธร|ยะลา|ร้อยเอ็ด|ระนอง|ระยอง|ราชบุรี|ลพบุรี|ลำปาง|ลำพูน|เลย|ศรีสะเกษ|สกลนคร|สงขลา|สตูล|สมุทรปราการ|สมุทรสงคราม|สมุทรสาคร|สระแก้ว|สระบุรี|สิงห์บุรี|สุโขทัย|สุพรรณบุรี|สุราษฎร์ธานี|สุรินทร์|หนองคาย|หนองบัวลำภู|อ่างทอง|อำนาจเจริญ|อุดรธานี|อุตรดิตถ์|อุทัยธานี|อุบลราชธานี".split('|'))
ST_MAP={'ได้รับสินค้าแล้ว':'ส่งสำเร็จ','พัสดุตีกลับ':'ตีกลับ','ยกเลิกออเดอร์':'ยกเลิก'}
PAY_MAP={'COD':'เก็บเงินปลายทาง','TRANSFER':'โอนเงิน'}
PROBLEM_PRESET={'ลูกค้าปฎิเสธรับพัสดุ':'ลูกค้าปฏิเสธรับพัสดุ','ไม่สามารถติดต่อได้':'ไม่สามารถติดต่อได้'}
# --- การตัดสินใจของเจ้านาย ---
PHONE_FIX={('ม.ค.69',2106):'0875402122'}
DATE_FIX={('มี.ค.69',8898):datetime.datetime(2026,3,31,0,0)}
TRK_FIX={('ม.ค.69',5761):'7228019382397976', ('ม.ค.69',6239):'7227015808026696'}
SPLIT={'7228011328095396':(3590,390),'7128021460228426':(1790,680),'7128028156621596':(0,1980),
       '7227022292574466':(690,1790),'7027039309329716':(1790,490),'7027056638395416':(1590,680),
       '7127055316617736':(2590,880),'GOSH0622950532':(1790,1790)}
# --- โหลดรหัสเซลที่เจ้านายกรอก + รายการตัวแทน ---
fx=openpyxl.load_workbook(FIX, data_only=True)
SELLER_FIX={}; SELLER_BLANK=set()
ws=fx['รหัสพนักงาน']; hdr=[c.value for c in ws[1]]; ci=hdr.index('รหัสที่ใช้นำเข้า')+1
for r in range(2, ws.max_row+1):
    k=(ws.cell(r,1).value, int(ws.cell(r,2).value)); v=ws.cell(r,ci).value
    if v and str(v).strip(): SELLER_FIX[k]=str(v).strip()
    else: SELLER_BLANK.add(k)
AGENTS=set()
ws=fx['ตัวแทน-คัดเพิ่ม']
for r in range(2, ws.max_row+1): AGENTS.add((ws.cell(r,1).value, int(ws.cell(r,2).value)))
def parse_dt(v, key):
    if key in DATE_FIX: return DATE_FIX[key]
    if isinstance(v, datetime.datetime): return v
    m=re.match(r'^\s*(\d{1,2})/(\d{1,2})/(\d{4})\s*,?\s*(\d{1,2}):(\d{2})\s*$', str(v or ''))
    if m: return datetime.datetime(int(m[3]),int(m[2]),int(m[1]),int(m[4]),int(m[5]))
    raise ValueError(f"วันที่อ่านไม่ได้ {key}: {v!r}")
def split_addr(a):
    a=re.sub(r'\s+',' ',str(a or '')).strip()
    m=re.search(r'^(.*?)\s+(\d{5})\s*$', a)
    body, zc = m.group(1).strip(), m.group(2)
    t=body.split(' '); pi=None
    for i in range(len(t)-1,-1,-1):
        p=re.sub(r'^(จ\.|จังหวัด)','',t[i]).strip()
        p='กรุงเทพมหานคร' if p in ('กทม','กทม.') else p
        if p in PROV: pi=i; prov=p; break
    rest=t[:pi]
    return (' '.join(rest[:-2]).strip() or None, rest[-2], rest[-1], prov, zc)
def brand_of(names, agent):
    if agent or ('เอกสาร' in names): return 'ตัวแทน'
    bs={('HOPEFUL' if n in HOPEFUL else 'แบรนด์อื่นๆ') for n in names}
    return bs.pop() if len(bs)==1 else None
COLS=['mo','src_row','order_no','ordered_at','customer_name','phone','addr_detail','subdistrict',
      'district','province','postal_code','seller_code','carrier','tracking_no','total_sales',
      'payment_method','payment_status','delivery_status','note','return_reason','status_detail','brand_name','items']
def run(mo):
    sheet, ymm = MONTHS[mo]
    wb=openpyxl.load_workbook(SRC, data_only=True); w=wb[sheet]
    out=[]; seq=0; stats=collections.Counter()
    for r in range(2, w.max_row+1):
        v=[w.cell(r,c).value for c in range(1,14)]
        if not any(v[:5]): continue
        key=(mo,r)
        dt=parse_dt(v[0], key)
        ph=PHONE_FIX.get(key) or re.sub(r'\D','',str(v[1] or ''))
        assert re.fullmatch(r'[0-9]{8,15}', ph), f"เบอร์ผิด {key}: {ph!r}"
        name=str(v[2] or '').strip()
        ad, sub, dis, prov, zc = split_addr(v[3])
        items=[(m.group(1).strip(), int(m.group(2))) for m in re.finditer(r'\[\s*(.+?)\s*\*\s*(\d+)\s*\]', str(v[4] or ''), re.S)]
        has_doc = any(n=='เอกสาร' for n,_ in items)
        agent = (key in AGENTS) or has_doc
        # tracking
        trk = TRK_FIX.get(key) or str(v[5] or '').strip()
        note=None
        if not re.fullmatch(r'[A-Za-z0-9\-]{8,}', trk):
            note = f"ช่องเลขพัสดุเดิมระบุว่า: {trk}"; trk=None; stats['trk_null']+=1
        carrier=str(v[6] or '').strip() or None
        pm = PAY_MAP.get(str(v[7] or '').strip())
        assert pm, f"การชำระเงินแปลกที่ {key}: {v[7]!r}"
        st_raw=str(v[9] or '').strip()
        ds = ST_MAP.get(st_raw, 'มีปัญหา')
        sd = None
        if ds=='มีปัญหา':
            sd = PROBLEM_PRESET.get(st_raw, st_raw); stats['problem']+=1
        # return_reason
        rr=None
        if ds=='ตีกลับ':
            why=str(v[11] or '').strip().replace('>>','').strip()
            extra=str(v[12] or '').strip()
            if why: rr=why
            if extra: sd=extra
        # seller
        if agent: sc=None
        elif key in SELLER_FIX: sc=SELLER_FIX[key]
        elif key in SELLER_BLANK: sc=None
        else:
            m=re.search(r'([A-Za-z]{1,4}\d{1,4})\s*$', name); sc=m.group(1) if m else None
        # payment_status
        if agent: ps='ไม่ใช่งานขาย'
        elif st_raw=='รับเงินเคลมแล้ว': ps='ชำระแล้ว'
        elif ds=='ส่งสำเร็จ': ps='ชำระแล้ว'
        else: ps='รอชำระ'
        amt=int(v[8] or 0)
        # ปนแบรนด์ → แยก 2 ใบ
        if trk in SPLIT and not agent:
            a,b=SPLIT[trk]
            legs=[('HOPEFUL', trk, a, [(n,q) for n,q in items if n in HOPEFUL]),
                  ('แบรนด์อื่นๆ', trk+'-B', b, [(n,q) for n,q in items if n not in HOPEFUL])]
            stats['split']+=1
        else:
            its = items + ([('เอกสาร',1)] if (agent and not has_doc) else [])
            bn = brand_of([n for n,_ in its], agent)
            assert bn, f"ปนแบรนด์ที่ไม่ได้วางแผน {key}: {items}"
            legs=[(bn, trk, amt, its)]
        for bn, tk, am, its in legs:
            seq+=1
            out.append(dict(mo=mo, src_row=r, order_no=f"HIST{ymm}-{seq:05d}",
                ordered_at=dt.strftime('%Y-%m-%d %H:%M:%S+07'), customer_name=name, phone=ph,
                addr_detail=ad, subdistrict=sub, district=dis, province=prov, postal_code=zc,
                seller_code=sc, carrier=carrier, tracking_no=tk, total_sales=am,
                payment_method=pm, payment_status=ps, delivery_status=ds, note=note,
                return_reason=rr, status_detail=sd, brand_name=bn,
                items=json.dumps([{"name":n,"qty":q} for n,q in its], ensure_ascii=False)))
            stats['rows']+=1; stats['items']+=len(its); stats['sales']+=am
    path=f"/tmp/hist/{ymm}.tsv"
    def esc(x):
        if x is None: return "\\N"
        return (str(x).replace("\\","\\\\").replace("\t","\\t")
                .replace("\n","\\n").replace("\r","\\r"))
    with open(path,"w") as f:
        for o in out:
            f.write("\t".join(esc(o[k]) for k in COLS) + "\n")
    print(f"{mo}: {stats['rows']:,} ออเดอร์ · {stats['items']:,} รายการ · {stats['sales']:,} บาท "
          f"· แยก {stats['split']} · trk ว่าง {stats['trk_null']} · มีปัญหา {stats['problem']} → {path}")
    return stats
if __name__=='__main__':
    for mo in (sys.argv[1:] or list(MONTHS)): run(mo)
